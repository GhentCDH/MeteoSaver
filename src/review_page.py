#!/usr/bin/env python
"""
Review page of one sheet: the cells of the table, column by column, with the crop that was read, the text as read
by GLM-OCR, and the value after QA/QC next to it (and the manual transcription, where there is one).

Written by main.py for the 'review' stage, which the dashboard (src/tui.py) shows as REVIEW PAGE and opens in the
browser. The page is one static HTML file with its crops in an img/ folder next to it: no server is needed.
"""
import json
import os

import cv2
import openpyxl

# Rows and columns of the pre-QA/QC workbook that hold the data (after 3 header rows, up to 43 days and totals)
FIRST_ROW, LAST_ROW, FIRST_COL, LAST_COL = 4, 46, 3, 26


def _grid(path, sheet=None):
    """ {(row, column): value} of the filled cells of a workbook; empty if there is no such file. """
    if not path or not os.path.exists(path):
        return {}
    workbook = openpyxl.load_workbook(path, data_only=True)
    worksheet = workbook[sheet] if sheet else workbook.active
    return {(c.row, c.column): c.value for row in worksheet.iter_rows() for c in row if c.value not in (None, '')}


def _number(value):
    try:
        return float(str(value).strip().replace(',', '.').replace('\u2212', '-'))
    except ValueError:
        return None


def _same(a, b):
    """ Whether two readings of a cell agree: None if one is missing. Numbers agree when equal, or when one is the other in tenths ('317' = 31.7). """
    if a in (None, '') or b in (None, ''):
        return None
    x, y = _number(a), _number(b)
    if x is None or y is None:
        return str(a).strip() == str(b).strip()
    return any(abs(candidate - y) < 0.051 for candidate in (x, x / 10))


def _show(value):
    if value is None:
        return ''
    if isinstance(value, float) and value == int(value):
        return str(int(value))
    return str(value)


def build_review_page(station, month_filename, clipped_table, glm_output_xlsx, post_QA_QC_xlsx, output_dir,
                      manual_xlsx=None, earlier_xlsx=None, model_note=''):
    '''
    Writes the review page of one sheet.

    Parameters
    --------------
    station, month_filename : str
        Station and image file name of the sheet.
    clipped_table : numpy.ndarray
        The clip of the original table (detected_table_cells[3] of table_and_cell_detection): the image GLM-OCR reads cells from.
    glm_output_xlsx : str
        The workbook saved by transcription() with 'GLM-OCR': the text per cell, and the bounding box of each cell.
    post_QA_QC_xlsx : str
        The post-QA/QC workbook of the sheet.
    output_dir : str
        Directory of the page: index.html, and img/ with the crops.
    manual_xlsx : str, optional
        Manual transcription of the sheet (temperature columns), to check the readings against.
    earlier_xlsx : str, optional
        Post-QA/QC workbook of an earlier run, to show next to this one.
    model_note : str
        Short description of the settings, shown under the title.

    Returns
    --------------
    page_path : str
        Path to index.html.
    '''
    glm_text = _grid(glm_output_xlsx, 'GLM_OCR_text')    # as read: with the decimal comma, before only the digits are kept
    corrected = _grid(post_QA_QC_xlsx)
    earlier = _grid(earlier_xlsx)
    manual = _grid(manual_xlsx)

    image_dir = os.path.join(output_dir, 'img')
    os.makedirs(image_dir, exist_ok=True)
    boxes = {}
    cells_sheet = openpyxl.load_workbook(glm_output_xlsx)['Detected_cells']
    for _cell, row, column, x, y, w, h, *_rest in cells_sheet.iter_rows(min_row=2, values_only=True):
        boxes[(row, column)] = (x, y, w, h)
    for (row, column), (x, y, w, h) in boxes.items():
        crop = clipped_table[y:y + h, x:x + w]
        if crop.size:
            cv2.imwrite(os.path.join(image_dir, f'{row}_{column}.png'), crop)

    rows = [{'r': r, 'label': _show(corrected.get((r, 2)) or earlier.get((r, 2)) or '')} for r in range(FIRST_ROW, LAST_ROW + 1)]
    columns = []
    for c in range(FIRST_COL, LAST_COL + 1):
        head = _show(corrected.get((3, c)) or corrected.get((2, c)) or corrected.get((1, c)) or f'col {c}')
        cells = []
        tally = {key: [0, 0] for key in ('raw_man', 'cor_man', 'prev_man', 'raw_cor')}
        for r in range(FIRST_ROW, LAST_ROW + 1):
            raw, cor, old, man = glm_text.get((r, c)), corrected.get((r, c)), earlier.get((r, c)), manual.get((r, c))
            status = {'rm': _same(raw, man), 'cm': _same(cor, man), 'pm': _same(old, man), 'rc': _same(raw, cor)}
            for key, flag in (('raw_man', 'rm'), ('cor_man', 'cm'), ('prev_man', 'pm'), ('raw_cor', 'rc')):
                if status[flag] is not None:
                    tally[key][0] += bool(status[flag])
                    tally[key][1] += 1
            image = f'img/{r}_{c}.png' if os.path.exists(os.path.join(image_dir, f'{r}_{c}.png')) else ''
            cells.append({'img': image, 'raw': _show(raw), 'cor': _show(cor), 'prev': _show(old), 'man': _show(man), **status})
        percent = lambda t: round(100 * t[0] / t[1], 1) if t[1] else None
        columns.append({'c': c, 'head': head, 'cells': cells, 'n_man': tally['raw_man'][1], **{k: percent(v) for k, v in tally.items()}})

    summary = {key: [sum(1 for col in columns for cell in col['cells'] if cell[flag] is True),
                     sum(1 for col in columns for cell in col['cells'] if cell[flag] is not None)]
               for key, flag in (('raw_man', 'rm'), ('cor_man', 'cm'), ('prev_man', 'pm'), ('raw_cor', 'rc'))}
    data = {'file': month_filename, 'station': station, 'rows': rows, 'columns': columns, 'summary': summary,
            'model': model_note or 'GLM-OCR', 'has_manual': bool(manual), 'has_prev': bool(earlier)}

    page_path = os.path.join(output_dir, 'index.html')
    with open(page_path, 'w', encoding='utf-8') as f:
        f.write(PAGE_TEMPLATE.replace('__TITLE__', month_filename).replace('__DATA__', json.dumps(data)))
    return page_path


PAGE_TEMPLATE = """<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>__TITLE__</title><style>
:root{--bg:#fff;--fg:#1c1c1e;--mut:#6b7280;--line:#e5e7eb;--card:#f8f9fb;--ok:#e6f4ea;--okfg:#14532d;--bad:#fde8e8;--badfg:#9b1c1c;--warn:#fef3c7;--warnfg:#92400e;--acc:#2563eb}
@media (prefers-color-scheme:dark){:root{--bg:#15171a;--fg:#e6e7ea;--mut:#9aa0aa;--line:#2b2f36;--card:#1d2025;--ok:#15331f;--okfg:#86efac;--bad:#3b1818;--badfg:#fca5a5;--warn:#3a2d0a;--warnfg:#fcd34d;--acc:#60a5fa}}
*{box-sizing:border-box}body{margin:0;background:var(--bg);color:var(--fg);font:14px/1.45 system-ui,sans-serif}
header{padding:14px 16px 8px;border-bottom:1px solid var(--line)}h1{font-size:18px;margin:0}header p{margin:2px 0 0;color:var(--mut);font-size:13px}
a{color:var(--acc)}nav{display:flex;gap:6px;overflow-x:auto;padding:10px 16px;border-bottom:1px solid var(--line);position:sticky;top:0;background:var(--bg);z-index:3}
nav button{flex:0 0 auto;border:1px solid var(--line);background:var(--card);color:var(--fg);border-radius:8px;padding:6px 10px;cursor:pointer;font:inherit;text-align:left}
nav button[aria-selected=true]{border-color:var(--acc);box-shadow:0 0 0 1px var(--acc)}nav small{display:block;color:var(--mut)}
.bar{display:flex;gap:16px;align-items:center;padding:10px 16px;flex-wrap:wrap;color:var(--mut)}.bar b{color:var(--fg)}
table{border-collapse:collapse;margin:0 16px 40px;width:calc(100% - 32px);max-width:900px}
th{position:sticky;top:57px;background:var(--bg);text-align:left;color:var(--mut);font-size:12px;padding:6px 8px;border-bottom:1px solid var(--line)}
td{padding:4px 8px;border-bottom:1px solid var(--line);vertical-align:middle;font-variant-numeric:tabular-nums}
td.lab{color:var(--mut);width:56px}td img{display:block;max-height:56px;max-width:220px;background:#fff;border:1px solid var(--line)}
.ok{background:var(--ok);color:var(--okfg)}.bad{background:var(--bad);color:var(--badfg)}.warn{background:var(--warn);color:var(--warnfg)}.none{color:var(--mut)}
.tag{font-size:11px;margin-left:6px;opacity:.8}.legend{padding:0 16px;color:var(--mut);font-size:12px}
@media (max-width:640px){th{top:96px}}
</style></head><body>
<header><h1 id="t"></h1><p id="sub"></p></header>
<nav id="tabs" role="tablist"></nav>
<div class="bar"><span id="stat"></span><label><input type="checkbox" id="diff"> differences only (QA/QC changed it, or it differs from manual)</label></div>
<div class="legend">Where a manual value exists (temperature columns) a cell is green if it agrees and red if not. Elsewhere the GLM cell is green when QA/QC
kept the reading and amber when QA/QC changed it. "Earlier Tesseract pipeline" is the stored results/02 output, produced with the old row grouping,
so on some sheets its rows can sit one slot off from the crops.</div>
<table><thead><tr><th>Row</th><th>Crop</th><th>GLM-OCR (raw)</th><th>After QA/QC</th><th id="th-prev">Earlier Tesseract pipeline</th><th id="th-man">Manual</th></tr></thead><tbody id="body"></tbody></table>
<script>
const D = __DATA__;
const $ = (id) => document.getElementById(id);
let cur = Math.max(0, parseInt((location.hash || "").slice(1)) || 0);
const pct = (v) => v === null ? "\u2013" : v + "%";
const f = (a) => a[1] ? (100 * a[0] / a[1]).toFixed(1) + "% (" + a[0] + "/" + a[1] + ")" : "\u2013";
const s = D.summary;
$("t").textContent = "Sheet " + D.file;
if (!D.has_prev) $("th-prev").remove();
if (!D.has_manual) $("th-man").remove();
$("sub").textContent = D.model + (D.has_manual ? " | raw = manual: " + f(s.raw_man) + " | after QA/QC = manual: " + f(s.cor_man) : "") +
  (D.has_prev && D.has_manual ? " | earlier pipeline = manual: " + f(s.prev_man) : "") + " | QA/QC kept GLM's reading: " + f(s.raw_cor);
const tabs = $("tabs");
D.columns.forEach((c, i) => {
  const b = document.createElement("button");
  b.setAttribute("role", "tab");
  b.innerHTML = "<b></b><small></small>";
  b.firstChild.textContent = c.head;
  b.lastChild.textContent = "col " + c.c + " | " + (c.n_man ? "raw=man " + pct(c.raw_man) : "kept " + pct(c.raw_cor));
  b.onclick = () => { cur = i; render(); };
  tabs.appendChild(b);
});
function cell(text, cls, tag) {
  const td = document.createElement("td");
  td.className = cls || "none";
  td.textContent = text === "" || text === null ? "\u2013" : text;
  if (tag) { const e = document.createElement("span"); e.className = "tag"; e.textContent = tag; td.appendChild(e); }
  return td;
}
const okbad = (v) => v === true ? "ok" : v === false ? "bad" : "none";
function render() {
  location.hash = cur;
  [...tabs.children].forEach((b, i) => b.setAttribute("aria-selected", i === cur));
  tabs.children[cur].scrollIntoView({inline: "center", block: "nearest"});
  const col = D.columns[cur], only = $("diff").checked, body = $("body");
  body.textContent = "";
  let shown = 0;
  col.cells.forEach((c, i) => {
    const hasMan = c.man !== "";
    const changed = c.rc === false;
    const wrong = hasMan && (c.rm === false || c.cm === false);
    if (only && !(changed || wrong)) return;
    shown++;
    const tr = document.createElement("tr");
    const lab = document.createElement("td"); lab.className = "lab"; lab.textContent = D.rows[i].label || D.rows[i].r; tr.appendChild(lab);
    const im = document.createElement("td");
    if (c.img) { const e = document.createElement("img"); e.src = c.img; e.loading = "lazy"; e.alt = ""; im.appendChild(e); } else { im.className = "none"; im.textContent = "\u2013"; }
    tr.appendChild(im);
    tr.appendChild(cell(c.raw, hasMan ? okbad(c.rm) : (changed ? "warn" : (c.rc === true ? "ok" : "none"))));
    tr.appendChild(cell(c.cor, hasMan ? okbad(c.cm) : (changed ? "warn" : "none"), changed ? "changed" : ""));
    if (D.has_prev) tr.appendChild(cell(c.prev, hasMan ? okbad(c.pm) : "none"));
    if (D.has_manual) tr.appendChild(cell(c.man, "none"));
    body.appendChild(tr);
  });
  $("stat").innerHTML = "<b>" + col.head + "</b> (col " + col.c + "): " +
    (col.n_man ? "raw = manual " + pct(col.raw_man) + " | after QA/QC = manual " + pct(col.cor_man) + " | earlier pipeline = manual " + pct(col.prev_man) + " | " : "") +
    "QA/QC kept GLM's reading " + pct(col.raw_cor) + " | showing " + shown + " rows";
}
$("diff").onchange = render;
document.addEventListener("keydown", (e) => {
  if (e.key === "ArrowRight") { cur = Math.min(D.columns.length - 1, cur + 1); render(); }
  if (e.key === "ArrowLeft") { cur = Math.max(0, cur - 1); render(); }
});
render();
</script></body></html>"""
