"""Cell-by-cell report for compare_sheet.py output (run with .venv-easyocr python)."""
import collections, json, sys
from pathlib import Path

import openpyxl

REPO = Path(__file__).resolve().parents[2]
FILE = sys.argv[1] if len(sys.argv) > 1 else "508_198802_SF.JPG"
STATION = FILE.split("_")[0]
E = REPO / "experiments/glm_ocr/sheet_runs" / FILE


def grid(p):
    ws = openpyxl.load_workbook(p).active
    return {(c.row, c.column): ("" if c.value is None else str(c.value).strip()) for r in ws.iter_rows() for c in r}


b = grid(E / f"baseline/pre_QA_QC/{FILE}_pre_QA_QC.xlsx")
g = grid(E / f"{sys.argv[2] if len(sys.argv) > 2 else 'glm'}/pre_QA_QC/{FILE}_pre_QA_QC.xlsx")
o = grid(REPO / f"results/01_pre_QA_QC_transcribed_hydroclimate_data/{STATION}/{FILE}_pre_QA_QC.xlsx")
data = [(r, c) for r in range(4, 50) for c in range(3, 27)]
nz = lambda d: sum(1 for k in data if d.get(k))
print("non-empty data cells: baseline", nz(b), "| glm", nz(g), "| existing results/01", nz(o), "| of", len(data))
print("baseline reproduces existing results/01:", sum(b.get(k, "") == o.get(k, "") for k in data), "/", len(data))
both = [k for k in data if b.get(k) or g.get(k)]
print("cells filled in either:", len(both), "| identical:", sum(b.get(k, "") == g.get(k, "") for k in both))
for k in [k for k in both if b.get(k, "") != g.get(k, "")][:30]:
    print(" ", k, "tess=", repr(b.get(k, "")), "glm=", repr(g.get(k, "")))
log = json.loads((E / f"{sys.argv[2] if len(sys.argv) > 2 else 'glm'}/glm_calls.json").read_text())
print("GLM calls", len(log), "| with '.':", sum("." in x["text"] for x in log), "| empty:", sum(not x["text"] for x in log), "| errors:", sum("error" in x for x in log))
print("length histogram:", sorted(collections.Counter(len(x["text"]) for x in log).items()))
m = openpyxl.load_workbook(REPO / f"results/06_manually_transcribed_hydroclimate_data/{STATION}/{FILE}_manually_entered_temperatures.xlsx").active
print("manual file dims", m.max_row, m.max_column)
for r in m.iter_rows(min_row=1, max_row=6, values_only=True):
    print(" ", r[:12])

# ---- Score against the manual temperature transcription (cols 4-6: max, min, (M+m)/2) ----
def tenths(v):
    return None if v in (None, "") else str(round(float(v) * 10))


def matches(ocr: str, truth: str) -> bool:
    d = ocr.replace(",", "").replace(".", "")
    return bool(d) and truth in (d, d + "0")  # '17' is accepted for 17.0 -> '170', for both systems alike


rows = [(r, c) for r in range(4, 47) for c in (4, 5, 6) if tenths(m.cell(r, c).value)]
print(f"\nScored against manual transcription: {len(rows)} temperature cells (cols Max/Min/Mean)")
for name, d in [("existing results/01", o), ("baseline (Tesseract chain)", b), ("GLM-OCR full", g)]:
    ok = sum(matches(d.get(k, ""), tenths(m.cell(*k).value)) for k in rows)
    empty = sum(not d.get(k, "") for k in rows)
    print(f"  {name:<28} correct {ok:>4}/{len(rows)} = {ok / len(rows):.1%}   empty {empty}")
print("\nGLM errors in scored cells (first 20):")
shown = 0
for k in rows:
    t = tenths(m.cell(*k).value)
    if not matches(g.get(k, ""), t) and shown < 20:
        print(f"  r{k[0]}c{k[1]} truth={m.cell(*k).value} glm={g.get(k, '')!r} tess={b.get(k, '')!r}")
        shown += 1
