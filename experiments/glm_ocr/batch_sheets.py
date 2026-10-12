"""Run GLM (original colour ROIs, rescaled to 128 px) on every sheet that has a manual
temperature transcription and score it against the existing pipeline output.

Run with .venv-easyocr/bin/python. Reuses compare_sheet.py per sheet (one GLM worker each).
Usage: .venv-easyocr/bin/python batch_sheets.py [--rerun]
"""
from __future__ import annotations

import argparse
import subprocess
import sys
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
REPO = HERE.parents[1]
MANUAL = REPO / "results/06_manually_transcribed_hydroclimate_data"
LABEL = "glm_orig_h128"
ROW_GROUPING = "kmeans"


def grid(path: Path) -> dict:
    ws = openpyxl.load_workbook(path).active
    return {(c.row, c.column): ("" if c.value is None else str(c.value).strip()) for r in ws.iter_rows() for c in r}


def tenths(v):
    return None if v in (None, "") else str(round(float(v) * 10))


def matches(ocr: str, truth: str) -> bool:
    d = ocr.replace(",", "").replace(".", "")
    return bool(d) and truth in (d, d + "0")


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--rerun", action="store_true")
    ap.add_argument("--only", nargs="*")
    ap.add_argument("--row-grouping", choices=["kmeans", "columnwise"], default="kmeans")
    args = ap.parse_args()
    global LABEL
    if args.row_grouping == "columnwise":
        LABEL = "glm_orig_h128_colwise"

    sheets = sorted((p.parent.name, p.name.split("_manually")[0]) for p in MANUAL.glob("*/*_manually_entered_temperatures.xlsx"))
    totals = {"n": 0, "old": 0, "glm": 0}
    print(f"{'sheet':<24}{'cells':>6}{'existing':>10}{'GLM':>8}")
    for station, f in sheets:
        if args.only and f not in args.only:
            continue
        out = HERE / "sheet_runs" / f / LABEL / "pre_QA_QC" / f"{f}_pre_QA_QC.xlsx"
        if args.rerun or not out.exists():
            r = subprocess.run([sys.executable, str(HERE / "compare_sheet.py"), "--station", station, "--file", f,
                                "--skip-baseline", "--resize-h", "128", "--roi-source", "original", "--label", LABEL, "--row-grouping", args.row_grouping],
                               capture_output=True, text=True, cwd=HERE)
            if not out.exists():
                print(f"{f:<24} FAILED: {r.stderr.strip().splitlines()[-1] if r.stderr.strip() else r.stdout[-200:]}", flush=True)
                continue
        m = openpyxl.load_workbook(MANUAL / station / f"{f}_manually_entered_temperatures.xlsx").active
        old = grid(REPO / f"results/01_pre_QA_QC_transcribed_hydroclimate_data/{station}/{f}_pre_QA_QC.xlsx")
        glm = grid(out)
        cells = [(r, c) for r in range(4, 47) for c in (4, 5, 6) if tenths(m.cell(r, c).value)]
        ok_old = sum(matches(old.get(k, ""), tenths(m.cell(*k).value)) for k in cells)
        ok_glm = sum(matches(glm.get(k, ""), tenths(m.cell(*k).value)) for k in cells)
        totals["n"] += len(cells); totals["old"] += ok_old; totals["glm"] += ok_glm
        print(f"{f:<24}{len(cells):>6}{ok_old / len(cells):>10.1%}{ok_glm / len(cells):>8.1%}", flush=True)
    n = totals["n"]
    if n:
        print(f"{'ALL':<24}{n:>6}{totals['old'] / n:>10.1%}{totals['glm'] / n:>8.1%}")


if __name__ == "__main__":
    main()
