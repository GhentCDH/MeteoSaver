"""Static HTML viewer: one sheet, column by column, ROI crop next to the raw OCR text and the corrected value.

For every sheet that has a manual transcription this runs the production path of src/main.py once
(detection, transcription() with ocr_model 'GLM-OCR' as set in configuration.ini, qa_qc()) into
view/_work/<file>/, then joins per cell
  crop | GLM raw text (decimal comma kept) | after QA/QC (same run) | earlier Tesseract pipeline (results/02) | manual (results/06)
and writes view/<file>/index.html (+ img/) and view/index.html. Existing results/ are never written.

Run with .venv-easyocr/bin/python.   Usage: visualise_sheet.py [--only FILE ...]
Open view/index.html in a browser; nothing is served or published.
"""
from __future__ import annotations

import argparse
import configparser
import json
import os
import shutil
from pathlib import Path

import openpyxl

import compare_sheet as cs

REPO, HERE = cs.REPO, cs.HERE
VIEW = HERE / "view"
FIRST_ROW, LAST_ROW, FIRST_COL, LAST_COL = 4, 46, 3, 26
MODEL = HERE / "runs/full/merged"
DATA_ROOT = REPO / "data/00_post1960_DRC_hydroclimate_datasheet_images"  # not configuration.ini: full_datadir may point at one station


def grid(path: Path, sheet: str | None = None) -> dict[tuple[int, int], object]:
    if not path.exists():
        return {}
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet] if sheet else wb.active
    return {(c.row, c.column): c.value for row in ws.iter_rows() for c in row if c.value not in (None, "")}


def num(v) -> float | None:
    if v is None:
        return None
    try:
        return float(str(v).strip().replace(",", ".").replace("\u2212", "-"))
    except ValueError:
        return None


def same(a, b) -> bool | None:
    """Do two readings of a cell agree? Numeric: equal, or equal when one is the other in tenths ('317' = 31.7)."""
    if a in (None, "") or b in (None, ""):
        return None
    x, y = num(a), num(b)
    if x is None or y is None:
        return str(a).strip() == str(b).strip()
    return any(abs(c - y) < 0.051 for c in (x, x / 10))


def show(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    return str(v)


def run_native(station: str, f: str, cfg, work: Path):
    """The production path of src/main.py for one sheet: detection, transcription() with GLM-OCR, qa_qc()."""
    import cv2
    import image_preprocessing_module as prep
    import quality_assessment_and_quality_control as qq
    import table_and_cell_detection_model as det
    import transcription as tr

    t, q = cfg["TableAndCellDetection"], cfg["QAQC"]
    transient, pre, post, glm_dir = work / "transient", work / "pre_QA_QC", work / "post_QA_QC", work / "glm"
    for d in (transient / station, pre, post):
        d.mkdir(parents=True, exist_ok=True)
    gray, binar, orig = prep.image_preprocessing(str(DATA_ROOT / station / f))
    detected = det.table_and_cell_detection(
        gray, binar, orig, station, f, str(transient),
        clip_up=int(t["clip_up"]), clip_down=int(t["clip_down"]), clip_left=int(t["clip_left"]), clip_right=int(t["clip_right"]),
        max_table_width=int(t["max_table_width"]), max_table_height=int(t["max_table_height"]),
        min_cell_width_threshold=int(t["min_cell_width_threshold"]), max_cell_width_threshold=int(t["max_cell_width_threshold"]),
        min_cell_height_threshold=int(t["min_cell_height_threshold"]), max_cell_height_threshold=int(t["max_cell_height_threshold"]),
        space_height_threshold=int(t["space_height_threshold"]), space_width_threshold=int(t["space_width_threshold"]),
        max_cell_height_per_box=int(t["max_cell_height_per_box"]), no_of_rows=int(t["no_of_rows"]), no_of_columns=int(t["no_of_columns"]))
    clip = detected[3]  # original colour table clip: what GLM-OCR reads in transcription()
    table = tr.transcription(
        detected, "GLM-OCR", "tesseract", str(transient), str(pre), station, f,
        no_of_rows=int(t["no_of_rows"]), no_of_columns=int(t["no_of_columns"]),
        no_of_rows_including_headers=int(t["no_of_rows_including_headers"]),
        row_grouping=cfg["Transcription"].get("row_grouping", "kmeans"),
        glm_model=str(MODEL), glm_python=str(cs.GLM_PY), glm_resize_height=int(cfg["Transcription"].get("glm_resize_height", 128)),
        glm_batch_size=int(cfg["Transcription"].get("glm_batch_size", 64)), glm_output_dir_station=str(glm_dir / station))
    qq.qa_qc(
        table, station, str(transient), str(post), f,
        max_temperature_threshold=float(q["max_temperature_threshold"]), min_temperature_threshold=float(q["min_temperature_threshold"]),
        decimal_places=int(q["decimal_places"]), uncertainty_margin=float(q["uncertainty_margin"]), header_rows=int(q["header_rows"]),
        multi_day_totals=cfg.getboolean("QAQC", "multi_day_totals"), multi_day_averages=cfg.getboolean("QAQC", "multi_day_averages"),
        max_days_for_multi_day_total=int(q["max_days_for_multi_day_total"]),
        multi_day_totals_rows=list(map(int, q["multi_day_totals_rows"].split(","))),
        final_totals_rows=list(map(int, q["final_totals_rows"].split(","))),
        excluded_rows=list(map(int, q["excluded_rows"].split(","))), excluded_columns=list(map(int, q["excluded_columns"].split(","))),
        daily_temperature_columns=q["daily_temperature_columns"].split(","),
        daily_temperature_columns_and_diurnal_temperature_range=q["daily_temperature_columns_and_diurnal_temperature_range"].split(","),
        daily_precipitation_column=q["daily_precipitation_column"].split(","),
        dry_and_wet_bulb_temperature_columns=q["dry_and_wet_bulb_temperature_columns"].split(","))
    return clip, glm_dir / station / f"{f}_GLM_OCR_output.xlsx", post / f"{f}_post_QA_QC.xlsx", cv2


def build_sheet(station: str, f: str, cfg) -> dict:
    work = VIEW / "_work" / f
    out = VIEW / f
    img_dir = out / "img"
    if img_dir.exists():
        shutil.rmtree(img_dir)
    img_dir.mkdir(parents=True, exist_ok=True)
    clip, glm_xlsx, post_xlsx, cv2 = run_native(station, f, cfg, work)

    glm_text = grid(glm_xlsx, "GLM_OCR_text")       # raw GLM text per cell, decimal comma kept
    corrected = grid(post_xlsx)                      # after QA/QC, same run
    prev = grid(REPO / f"results/02_post_QA_QC_transcribed_hydroclimate_data/{station}/{f}_post_QA_QC.xlsx")  # earlier Tesseract pipeline
    manual = grid(REPO / f"results/06_manually_transcribed_hydroclimate_data/{station}/{f}_manually_entered_temperatures.xlsx")
    boxes = {}
    ws = openpyxl.load_workbook(glm_xlsx)["Detected_cells"]
    for ref, r, c, x, y, w, h, *_ in ws.iter_rows(min_row=2, values_only=True):
        boxes[(r, c)] = (x, y, w, h)
    for (r, c), (x, y, w, h) in boxes.items():
        cv2.imwrite(str(img_dir / f"{r}_{c}.png"), clip[y:y + h, x:x + w])
    print(f"{f}: {len(boxes)} cells read, crops saved", flush=True)

    rows = [{"r": r, "label": show(corrected.get((r, 2)) or prev.get((r, 2)) or "")} for r in range(FIRST_ROW, LAST_ROW + 1)]
    columns = []
    for c in range(FIRST_COL, LAST_COL + 1):
        head = show(corrected.get((3, c)) or corrected.get((2, c)) or corrected.get((1, c)) or f"col {c}")
        cells = []
        tally = {k: [0, 0] for k in ("raw_man", "cor_man", "prev_man", "raw_cor")}
        for r in range(FIRST_ROW, LAST_ROW + 1):
            raw, cor, old, man = glm_text.get((r, c)), corrected.get((r, c)), prev.get((r, c)), manual.get((r, c))
            st = {"rm": same(raw, man), "cm": same(cor, man), "pm": same(old, man), "rc": same(raw, cor)}
            for k, key in (("raw_man", "rm"), ("cor_man", "cm"), ("prev_man", "pm"), ("raw_cor", "rc")):
                if st[key] is not None:
                    tally[k][0] += bool(st[key])
                    tally[k][1] += 1
            has_img = (r, c) in boxes
            cells.append({"img": f"img/{r}_{c}.png" if has_img else "", "raw": show(raw), "cor": show(cor), "prev": show(old), "man": show(man), **st})
        pct = lambda a: round(100 * a[0] / a[1], 1) if a[1] else None
        columns.append({"c": c, "head": head, "cells": cells, "n_man": tally["raw_man"][1],
                        **{k: pct(v) for k, v in tally.items()}})

    summary = {k: [sum(1 for col in columns for cell in col["cells"] if cell[s] is True),
                   sum(1 for col in columns for cell in col["cells"] if cell[s] is not None)]
               for k, s in (("raw_man", "rm"), ("cor_man", "cm"), ("prev_man", "pm"), ("raw_cor", "rc"))}
    data = {"file": f, "station": station, "rows": rows, "columns": columns, "summary": summary,
            "model": f"GLM-OCR ({MODEL.parent.name}), {cfg['Transcription'].get('row_grouping', 'kmeans')} row grouping"}
    out.mkdir(parents=True, exist_ok=True)
    (out / "index.html").write_text(SHEET_HTML.replace("__TITLE__", f).replace("__DATA__", json.dumps(data)), encoding="utf-8")
    return {"file": f, "summary": summary, "cells": len(boxes)}


def write_index(items: list[dict]) -> None:
    def pct(a):
        return f"{100 * a[0] / a[1]:.1f}%" if a[1] else "\u2013"
    rows = "\n".join(
        f'<tr><td><a href="{i["file"]}/index.html">{i["file"]}</a></td><td>{pct(i["summary"]["raw_man"])}</td>'
        f'<td>{pct(i["summary"]["cor_man"])}</td><td>{pct(i["summary"]["prev_man"])}</td><td>{pct(i["summary"]["raw_cor"])}</td>'
        f'<td>{i["cells"]}</td></tr>' for i in items)
    (VIEW / "index.html").write_text(INDEX_HTML.replace("__ROWS__", rows), encoding="utf-8")


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--only", nargs="*")
    args = ap.parse_args()

    cfg = configparser.ConfigParser()
    cfg.read(REPO / "configuration.ini")
    os.environ["TESSDATA_PREFIX"] = str(REPO / "OCR_HTR_models")
    VIEW.mkdir(exist_ok=True)
    sheets = sorted((p.parent.name, p.name.split("_manually")[0]) for p in
                    (REPO / "results/06_manually_transcribed_hydroclimate_data").glob("*/*_manually_entered_temperatures.xlsx"))
    items = []
    for station, f in sheets:
        if args.only and f not in args.only:
            continue
        items.append(build_sheet(station, f, cfg))
        write_index(items)
    print(f"open {VIEW / 'index.html'}")


INDEX_HTML = """<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>OCR check sheets</title><style>
:root{--bg:#fff;--fg:#1c1c1e;--mut:#6b7280;--line:#e5e7eb;--card:#f8f9fb}
@media (prefers-color-scheme:dark){:root{--bg:#15171a;--fg:#e6e7ea;--mut:#9aa0aa;--line:#2b2f36;--card:#1d2025}}
body{margin:0;padding:24px 16px;background:var(--bg);color:var(--fg);font:15px/1.5 system-ui,sans-serif}
main{max-width:860px;margin:0 auto}h1{font-size:22px;margin:0 0 4px}p{color:var(--mut);margin:0 0 18px}
table{border-collapse:collapse;width:100%}th,td{padding:8px 10px;border-bottom:1px solid var(--line);text-align:right}
th:first-child,td:first-child{text-align:left}th{color:var(--mut);font-weight:600;font-size:13px}a{color:inherit;font-weight:600}
</style></head><body><main><h1>OCR check: GLM-OCR vs corrected values</h1>
<p>Agreement per sheet, over cells where both values exist. The manual transcription covers the temperature columns only,
so the last-but-one column (QA/QC left GLM's reading unchanged) is the one that covers every column. Open a sheet to go through it
column by column.</p>
<table><thead><tr><th>Sheet</th><th>GLM raw = manual</th><th>GLM after QA/QC = manual</th><th>Earlier Tesseract pipeline = manual</th><th>QA/QC left GLM unchanged</th><th>Cells</th></tr></thead>
<tbody>__ROWS__</tbody></table></main></body></html>"""

SHEET_HTML = """<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>__TITLE__</title><style>
:root{--bg:#fff;--fg:#1c1c1e;--mut:#6b7280;--line:#e5e7eb;--card:#f8f9fb;--ok:#e6f4ea;--okfg:#14532d;--bad:#fde8e8;--badfg:#9b1c1c;--warn:#fef3c7;--warnfg:#92400e;--acc:#2563eb}
@media (prefers-color-scheme:dark){:root{--bg:#15171a;--fg:#e6e7ea;--mut:#9aa0aa;--line:#2b2f36;--card:#1d2025;--ok:#15331f;--okfg:#86efac;--bad:#3b1818;--badfg:#fca5a5;--warn:#3a2d0a;--warnfg:#fcd34d;--acc:#60a5fa}}
*{box-sizing:border-box}body{margin:0;background:var(--bg);color:var(--fg);font:14px/1.45 system-ui,sans-serif}
header{padding:14px 16px 8px;border-bottom:1px solid var(--line)}h1{font-size:18px;margin:0}header p{margin:2px 0 0;color:var(--mut);font-size:13px}
a{color:var(--acc)}nav{display:flex;gap:6px;overflow-x:auto;padding:10px 16px;border-bottom:1px solid var(--line);position:sticky;top:0;background:var(--bg);z-index:3}
nav button{flex:0 0 auto;border:1px solid var(--line);background:var(--card);color:var(--fg);border-radius:8px;padding:6px 10px;cursor:pointer;font:inherit;text-align:left}
nav button[aria-selected=true]{border-color:var(--acc);box-shadow:0 0 0 1px var(--acc)}nav small{display:block;color:var(--mut)}
.bar{display:flex;gap:16px;align-items:center;padding:10px 16px;flex-wrap:wrap;color:var(--mut)}.bar b{color:var(--fg)}
table{border-collapse:collapse;margin:0 16px 40px;width:calc(100% - 32px);max-width:900px}
th{position:sticky;top:57px;background:var(--bg);text-align:left;color:var(--mut);font-size:12px;padding:6px 8px;border-bottom:1px solid var(--line)}
td{padding:4px 8px;border-bottom:1px solid var(--line);vertical-align:middle;font-variant-numeric:tabular-nums}
td.lab{color:var(--mut);width:56px}td img{display:block;max-height:56px;max-width:220px;background:#fff;border:1px solid var(--line)}
.ok{background:var(--ok);color:var(--okfg)}.bad{background:var(--bad);color:var(--badfg)}.warn{background:var(--warn);color:var(--warnfg)}.none{color:var(--mut)}
.tag{font-size:11px;margin-left:6px;opacity:.8}.legend{padding:0 16px;color:var(--mut);font-size:12px}
@media (max-width:640px){th{top:96px}}
</style></head><body>
<header><h1 id="t"></h1><p id="sub"></p><p><a href="../index.html">All sheets</a></p></header>
<nav id="tabs" role="tablist"></nav>
<div class="bar"><span id="stat"></span><label><input type="checkbox" id="diff"> differences only (QA/QC changed it, or it differs from manual)</label></div>
<div class="legend">Where a manual value exists (temperature columns) a cell is green if it agrees and red if not. Elsewhere the GLM cell is green when QA/QC
kept the reading and amber when QA/QC changed it. "Earlier Tesseract pipeline" is the stored results/02 output, produced with the old row grouping,
so on some sheets its rows can sit one slot off from the crops.</div>
<table><thead><tr><th>Row</th><th>Crop</th><th>GLM-OCR (raw)</th><th>After QA/QC</th><th>Earlier Tesseract pipeline</th><th>Manual</th></tr></thead><tbody id="body"></tbody></table>
<script>
const D = __DATA__;
const $ = (id) => document.getElementById(id);
let cur = Math.max(0, parseInt((location.hash || "").slice(1)) || 0);
const pct = (v) => v === null ? "\u2013" : v + "%";
const f = (a) => a[1] ? (100 * a[0] / a[1]).toFixed(1) + "% (" + a[0] + "/" + a[1] + ")" : "\u2013";
const s = D.summary;
$("t").textContent = "Sheet " + D.file;
$("sub").textContent = D.model + " | raw = manual: " + f(s.raw_man) + " | after QA/QC = manual: " + f(s.cor_man) +
  " | earlier pipeline = manual: " + f(s.prev_man) + " | QA/QC kept GLM's reading: " + f(s.raw_cor);
const tabs = $("tabs");
D.columns.forEach((c, i) => {
  const b = document.createElement("button");
  b.setAttribute("role", "tab");
  b.innerHTML = "<b></b><small></small>";
  b.firstChild.textContent = c.head;
  b.lastChild.textContent = "col " + c.c + " | " + (c.n_man ? "raw=man " + pct(c.raw_man) : "kept " + pct(c.raw_cor));
  b.onclick = () => { cur = i; render(); };
  tabs.appendChild(b);
});
function cell(text, cls, tag) {
  const td = document.createElement("td");
  td.className = cls || "none";
  td.textContent = text === "" || text === null ? "\u2013" : text;
  if (tag) { const e = document.createElement("span"); e.className = "tag"; e.textContent = tag; td.appendChild(e); }
  return td;
}
const okbad = (v) => v === true ? "ok" : v === false ? "bad" : "none";
function render() {
  location.hash = cur;
  [...tabs.children].forEach((b, i) => b.setAttribute("aria-selected", i === cur));
  tabs.children[cur].scrollIntoView({inline: "center", block: "nearest"});
  const col = D.columns[cur], only = $("diff").checked, body = $("body");
  body.textContent = "";
  let shown = 0;
  col.cells.forEach((c, i) => {
    const hasMan = c.man !== "";
    const changed = c.rc === false;
    const wrong = hasMan && (c.rm === false || c.cm === false);
    if (only && !(changed || wrong)) return;
    shown++;
    const tr = document.createElement("tr");
    const lab = document.createElement("td"); lab.className = "lab"; lab.textContent = D.rows[i].label || D.rows[i].r; tr.appendChild(lab);
    const im = document.createElement("td");
    if (c.img) { const e = document.createElement("img"); e.src = c.img; e.loading = "lazy"; e.alt = ""; im.appendChild(e); } else { im.className = "none"; im.textContent = "\u2013"; }
    tr.appendChild(im);
    tr.appendChild(cell(c.raw, hasMan ? okbad(c.rm) : (changed ? "warn" : (c.rc === true ? "ok" : "none"))));
    tr.appendChild(cell(c.cor, hasMan ? okbad(c.cm) : (changed ? "warn" : "none"), changed ? "changed" : ""));
    tr.appendChild(cell(c.prev, hasMan ? okbad(c.pm) : "none"));
    tr.appendChild(cell(c.man, "none"));
    body.appendChild(tr);
  });
  $("stat").innerHTML = "<b>" + col.head + "</b> (col " + col.c + "): " +
    (col.n_man ? "raw = manual " + pct(col.raw_man) + " | after QA/QC = manual " + pct(col.cor_man) + " | earlier pipeline = manual " + pct(col.prev_man) + " | " : "") +
    "QA/QC kept GLM's reading " + pct(col.raw_cor) + " | showing " + shown + " rows";
}
$("diff").onchange = render;
document.addEventListener("keydown", (e) => {
  if (e.key === "ArrowRight") { cur = Math.min(D.columns.length - 1, cur + 1); render(); }
  if (e.key === "ArrowLeft") { cur = Math.max(0, cur - 1); render(); }
});
render();
</script></body></html>"""


if __name__ == "__main__":
    main()
