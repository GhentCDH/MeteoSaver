"""Export every detected cell box and its GLM-OCR text for a station's sheets into one Excel workbook.

Runs the real detection + transcription() with GLM in place of the Tesseract call (as compare_sheet.py
does: original colour ROIs, rescaled to 128 px, column-wise row grouping) and records, per OCR call, the
ROI box transcription() cut and the sheet cell it wrote the text to. src/ is not modified.

Per sheet the workbook gets two tabs:
  <yyyymm> transcription = the pre-QA/QC table as the pipeline writes it
  <yyyymm> detections    = one row per detected cell: sheet cell, box (px, in the clipped table image), GLM text

Run with .venv-easyocr/bin/python.   Usage: export_detections.py --station 203 [--out FILE]
"""
from __future__ import annotations

import argparse
import configparser
import os
import sys
import types
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font
from openpyxl.utils import column_index_from_string, get_column_letter

import compare_sheet as cs

LABEL = "glm_export"
COLUMNS = ["No de la pentade", "Date", "Bellani", "Max.", "Min.", "(M+m)/2", "Ampl.", "Min. gazon", "Piche Abri.", "Piche Ext.", "Pluies",
           "6h T", "6h T'a", "6h e.", "6h U", "6h ∆e", "15h T", "15h T'a", "15h e.", "15h U", "15h ∆e",
           "18h T", "18h T'a", "18h e.", "18h U", "18h ∆e", "Date"]


class Recorder:
    """Wraps the GLM client: notes the ROI box from transcription()'s frame, then the cell it is written to."""

    def __init__(self, client, tr):
        self.client, self.rows = client, []
        self._cell_ref = tr.calculate_cell_reference
        tr.calculate_cell_reference = self.cell_reference
        self._tr = tr

    def __call__(self, image_path, *a, **k) -> str:
        f = sys._getframe(1).f_locals
        text = self.client(image_path)
        self.rows.append({"method": f["method_name"], "x": f["x"], "y": f["y"], "w": f["w"], "h": f["h"],
                          "text": text.strip(), "raw": self.client.log[-1].get("raw", ""), "ref": None})
        return text

    def cell_reference(self, *a, **k):
        ref = self._cell_ref(*a, **k)
        self.rows[-1]["ref"] = ref
        return ref

    def restore(self):
        self._tr.calculate_cell_reference = self._cell_ref


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--station", required=True)
    ap.add_argument("--model", default=str(cs.HERE / "runs/full/merged"))
    ap.add_argument("--out")
    args = ap.parse_args()
    cfg = configparser.ConfigParser()
    cfg.read(cs.REPO / "configuration.ini")
    os.environ["TESSDATA_PREFIX"] = str(cs.REPO / "OCR_HTR_models")
    t = cfg["TableAndCellDetection"]
    header_rows = int(t["no_of_rows_including_headers"]) - int(t["no_of_rows"])

    out = Path(args.out) if args.out else cs.HERE / "sheet_runs" / f"{args.station}_GLM_detections_and_OCR.xlsx"
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    bold = Font(bold=True)
    tr = cs.load_transcription(False)
    client = cs.GlmClient(args.model, 128)
    try:
        for image in sorted((cs.REPO / cfg["Directories"]["full_datadir"] / args.station).glob("*.JPG")):
            ns = types.SimpleNamespace(station=args.station, file=image.name)
            rec = Recorder(client, tr)
            try:
                xlsx = cs.run_pipeline(LABEL, ns, cfg, patch_ocr=rec, roi_source="original", row_grouping="columnwise")
            finally:
                rec.restore()
            tag = image.name.split("_")[1]

            ws = wb.create_sheet(f"{tag} transcription")
            src = openpyxl.load_workbook(xlsx).active
            for row in src.iter_rows():
                for c in row:
                    ws.cell(c.row, c.column, c.value)
            for rng in src.merged_cells.ranges:
                ws.merge_cells(str(rng))
            for r in range(1, header_rows + 1):
                for c in ws[r]:
                    c.font, c.alignment = bold, Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.freeze_panes = ws.cell(header_rows + 1, 3)

            ws = wb.create_sheet(f"{tag} detections")
            ws.append(["id", "sheet cell", "sheet row", "sheet column", "variable", "x", "y", "w", "h", "GLM text", "GLM raw output", "in transcription tab"])
            for c in ws[1]:
                c.font = bold
            final = {(c.row, c.column): ("" if c.value is None else str(c.value).strip()) for row in src.iter_rows() for c in row}
            # Top and Midpoint passes use the same grouping here, so they see the same boxes; list one pass.
            for i, d in enumerate((d for d in rec.rows if d["method"] == "Top"), start=1):
                col = row = cell = var = kept = None
                if d["ref"]:
                    letters = d["ref"].rstrip("0123456789")
                    col = column_index_from_string(letters) + 2  # two columns inserted on the left
                    row = int(d["ref"][len(letters):]) + header_rows
                    cell, var = f"{get_column_letter(col)}{row}", COLUMNS[col - 1] if col <= len(COLUMNS) else None
                    kept = "yes" if final.get((row, col), "") == d["text"] else "no"
                ws.append([i, cell, row, col, var, d["x"], d["y"], d["w"], d["h"], d["text"], d["raw"], kept])
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for letter, width in zip("ABCDEFGHIJKL", (6, 11, 10, 13, 16, 7, 7, 6, 6, 12, 26, 20)):
                ws.column_dimensions[letter].width = width
            print(f"{image.name}: {ws.max_row - 1} detections", flush=True)
    finally:
        client.close()
    wb.save(out)
    print(out)


if __name__ == "__main__":
    main()
